from io import BytesIO
from pathlib import Path
import re
from docx import Document
from pypdf import PdfReader

def parse_resume(filename,data):
 ext=Path(filename or '').suffix.lower()
 if ext=='.pdf': text='\n'.join((p.extract_text() or '') for p in PdfReader(BytesIO(data)).pages)
 elif ext=='.docx':
  doc=Document(BytesIO(data));chunks=[p.text for p in doc.paragraphs if p.text.strip()]
  for table in doc.tables:
   for row in table.rows: chunks.append(' | '.join(c.text.strip() for c in row.cells))
  text='\n'.join(chunks)
 elif ext in {'.txt','.md'}: text=data.decode('utf-8',errors='ignore')
 else: raise ValueError('Upload a PDF, DOCX, TXT, or MD resume.')
 text=re.sub(r'\n{3,}','\n\n',text).strip()
 if len(text)<50: raise ValueError('The resume text could not be extracted. Use a text-based PDF or DOCX.')
 return text
